import os
import sys
import time
import re
import json
import sqlite3
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.common.exceptions import WebDriverException, TimeoutException
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo


class IMDBPageDownloader:
    def __init__(self, html_dir="imdb_pages"):
        self.html_dir = html_dir
        self.driver = None
        os.makedirs(self.html_dir, exist_ok=True)

    def _create_driver(self, headless=False):
        options = Options()
        if headless:
            options.add_argument("--headless=new")
        options.add_argument("--window-size=1920,1080")
        options.add_argument("--lang=en-US")
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-blink-features=AutomationControlled")
        options.add_experimental_option("excludeSwitches", ["enable-automation"])
        options.add_experimental_option("useAutomationExtension", False)
        options.add_argument(
            "user-agent=Mozilla/5.0 (Linux NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
        try:
            service = Service('setup your chrome')  
            driver = webdriver.Chrome(service=service, options=options)
            driver.execute_script(
                "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
            )
            return driver
        except Exception as e:
            print(f"Failed to initialize ChromeDriver: {e}")
            sys.exit(1)

    def start(self):
        self.driver = self._create_driver()

    def stop(self):
        if self.driver:
            self.driver.quit()

    def fetch(self, url, movie_id, retries=3):
        html_path = os.path.join(self.html_dir, f"{movie_id}.html")
        if os.path.exists(html_path):
            print(f"  [CACHE] {movie_id} already saved")
            return html_path
        for attempt in range(retries):
            try:
                self.driver.get(url)
                WebDriverWait(self.driver, 10).until(
                    EC.presence_of_element_located((By.CSS_SELECTOR, 'h1[data-testid="hero__pageTitle"]'))
                )
                time.sleep(2)
                with open(html_path, 'w', encoding='utf-8') as f:
                    f.write(self.driver.page_source)
                print(f"  [SAVED] {movie_id}")
                return html_path
            except (WebDriverException, TimeoutException) as e:
                print(f"  Attempt {attempt+1}/{retries} failed for {movie_id}: {e}")
                time.sleep(3)
                if attempt == retries - 1:
                    return None
                self.driver.quit()
                self.driver = self._create_driver()
        return None


class IMDBDataExtractor:
    def __init__(self, db_name="imdb_movies.db"):
        self.db_name = db_name
        self.conn = None
        self.person_cache = set()

    def connect(self):
        self.conn = sqlite3.connect(self.db_name)
        self._create_tables()

    def close(self):
        if self.conn:
            self.conn.close()

    def _create_tables(self):
        cursor = self.conn.cursor()
        cursor.executescript('''
            CREATE TABLE IF NOT EXISTS movies (
                id INTEGER PRIMARY KEY,
                title TEXT,
                year INTEGER,
                parental_guide TEXT,
                runtime INTEGER,
                genre TEXT
            );
            CREATE TABLE IF NOT EXISTS people (
                id INTEGER PRIMARY KEY,
                name TEXT
            );
            CREATE TABLE IF NOT EXISTS movie_directors (
                movie_id INTEGER,
                person_id INTEGER,
                FOREIGN KEY (movie_id) REFERENCES movies (id),
                FOREIGN KEY (person_id) REFERENCES people (id),
                PRIMARY KEY (movie_id, person_id)
            );
            CREATE TABLE IF NOT EXISTS movie_writers (
                movie_id INTEGER,
                person_id INTEGER,
                FOREIGN KEY (movie_id) REFERENCES movies (id),
                FOREIGN KEY (person_id) REFERENCES people (id),
                PRIMARY KEY (movie_id, person_id)
            );
            CREATE TABLE IF NOT EXISTS movie_stars (
                movie_id INTEGER,
                person_id INTEGER,
                FOREIGN KEY (movie_id) REFERENCES movies (id),
                FOREIGN KEY (person_id) REFERENCES people (id),
                PRIMARY KEY (movie_id, person_id)
            );
            CREATE TABLE IF NOT EXISTS box_office (
                movie_id INTEGER PRIMARY KEY,
                gross_us_canada TEXT,
                FOREIGN KEY (movie_id) REFERENCES movies (id)
            );
        ''')
        self.conn.commit()

    def _parse_runtime(self, text):
        if not text:
            return 0
        h = re.search(r'(\d+)h', text)
        m = re.search(r'(\d+)m', text)
        return (int(h.group(1)) if h else 0) * 60 + (int(m.group(1)) if m else 0)

    def _parse_year(self, text):
        if not text:
            return 0
        match = re.search(r'\d{4}', text)
        return int(match.group()) if match else 0

    def _extract_people(self, soup, keyword):
        people = []
        details_list = soup.select_one('ul[data-testid="title-pc-list"]')
        if not details_list:
            return people
        for item in details_list.select('li.ipc-metadata-list__item'):
            try:
                label = item.select_one('a.ipc-metadata-list-item__label')
                if not label:
                    continue
                if keyword not in label.get_text(strip=True).lower():
                    continue
                content = item.select_one('div.ipc-metadata-list-item__content-container')
                if not content:
                    continue
                for link in content.select('a[href*="/name/"]'):
                    href = link.get('href', '')
                    match = re.search(r'/name/nm(\d+)/', href)
                    if match:
                        people.append((match.group(1), link.get_text(strip=True)))
            except Exception:
                continue
        return people

    def _extract_gross(self, soup):
        gross = 'N/A'
        box_section = soup.select_one('section[data-testid="BoxOffice"]')
        if box_section:
            for item in box_section.select('li.ipc-metadata-list__item'):
                try:
                    label = item.select_one('span.ipc-metadata-list-item__label')
                    if not label:
                        continue
                    text = label.get_text(strip=True).lower()
                    if 'gross us & canada' in text or ('gross' in text and 'us' in text and 'canada' in text):
                        content = item.select_one('div.ipc-metadata-list-item__content-container')
                        if content:
                            raw = content.get_text(strip=True)
                            match = re.search(r'[\$]?[\d,]+', raw)
                            if match:
                                gross = match.group()
                                break
                except Exception:
                    continue
        if gross == 'N/A':
            details_section = soup.select_one('section[data-testid="title-details-section"]')
            if details_section:
                for item in details_section.select('li.ipc-metadata-list__item'):
                    try:
                        label = item.select_one('span.ipc-metadata-list-item__label')
                        if not label:
                            continue
                        text = label.get_text(strip=True).lower()
                        if 'gross' in text and ('usa' in text or 'canada' in text):
                            content = item.select_one('div.ipc-metadata-list-item__content-container')
                            if content:
                                raw = content.get_text(strip=True)
                                match = re.search(r'[\$]?[\d,]+', raw)
                                if match:
                                    gross = match.group()
                                    break
                    except Exception:
                        continue
        return gross

    def extract_from_html(self, html_path):
        with open(html_path, 'r', encoding='utf-8') as f:
            html = f.read()
        soup = BeautifulSoup(html, 'html.parser')
        details = {
            'title': 'N/A',
            'year': 0,
            'parental_guide': 'N/A',
            'runtime': 0,
            'genre': 'N/A',
            'directors': [],
            'writers': [],
            'stars': [],
            'gross': 'N/A'
        }
        script = soup.select_one('script[type="application/ld+json"]')
        if script:
            try:
                data = json.loads(script.string)
                if isinstance(data, list):
                    data = data[0]
                details['title'] = data.get('name', details['title'])
                if data.get('datePublished'):
                    details['year'] = self._parse_year(data['datePublished'])
                if data.get('duration'):
                    details['runtime'] = self._parse_runtime(data['duration'])
                if data.get('genre'):
                    details['genre'] = ', '.join(data['genre']) if isinstance(data['genre'], list) else data['genre']
                if data.get('director'):
                    dirs = []
                    for d in data['director']:
                        if isinstance(d, dict) and d.get('url'):
                            match = re.search(r'/name/nm(\d+)/', d['url'])
                            if match:
                                dirs.append((match.group(1), d.get('name', 'Unknown')))
                    details['directors'] = dirs
                if data.get('actor'):
                    actors = []
                    for a in data['actor'][:10]:
                        if isinstance(a, dict) and a.get('url'):
                            match = re.search(r'/name/nm(\d+)/', a['url'])
                            if match:
                                actors.append((match.group(1), a.get('name', 'Unknown')))
                    details['stars'] = actors
            except Exception:
                pass
        if details['title'] == 'N/A':
            title_elem = soup.select_one('h1[data-testid="hero__pageTitle"] span[data-testid="hero__primary-text"]')
            details['title'] = title_elem.get_text(strip=True) if title_elem else self._extract_text(soup, 'h1', 'N/A')
        if details['year'] == 0:
            year_elem = soup.select_one('a[href*="/releaseinfo"]')
            if year_elem:
                details['year'] = self._parse_year(year_elem.get_text(strip=True))
            else:
                details['year'] = self._parse_year(self._extract_text(soup, 'span[data-testid="title-year"]'))
        if details['parental_guide'] == 'N/A':
            details['parental_guide'] = self._extract_text(soup, 'span[data-testid="certificate"]')
        if details['runtime'] == 0:
            details['runtime'] = self._parse_runtime(self._extract_text(soup, 'span[data-testid="runtime"]'))
        if details['genre'] == 'N/A':
            genres = soup.select('div[data-testid="genres"] a')
            if genres:
                details['genre'] = ', '.join(g.get_text(strip=True) for g in genres)
            else:
                interests = soup.select('div[data-testid="interests"] a.ipc-chip')
                if interests:
                    details['genre'] = ', '.join(chip.get_text(strip=True) for chip in interests)
        if not details['directors']:
            details['directors'] = self._extract_people(soup, 'director')
        if not details['writers']:
            details['writers'] = self._extract_people(soup, 'writer')
        if not details['stars']:
            details['stars'] = self._extract_people(soup, 'star')
            if not details['stars']:
                cast_section = soup.select_one('section[data-testid="title-cast-section"]')
                if cast_section:
                    for item in cast_section.select('div[data-testid="title-cast-item"]')[:10]:
                        link = item.select_one('a[href*="/name/"]')
                        if link:
                            href = link.get('href', '')
                            match = re.search(r'/name/nm(\d+)/', href)
                            if match:
                                details['stars'].append((match.group(1), link.get_text(strip=True)))
        details['gross'] = self._extract_gross(soup)
        
        for key, default in [('title', 'N/A'), ('parental_guide', 'N/A'), ('genre', 'N/A'), ('gross', 'N/A')]:
            if details[key] is None:
                details[key] = default
        if details['year'] is None:
            details['year'] = 0
        if details['runtime'] is None:
            details['runtime'] = 0
        return details

    def _extract_text(self, soup, selector, default='N/A'):
        elem = soup.select_one(selector)
        return elem.get_text(strip=True) if elem else default

    def _insert_people(self, people):
        if not people:
            return
        cursor = self.conn.cursor()
        to_insert = [(pid, name) for pid, name in people if pid not in self.person_cache]
        if to_insert:
            cursor.executemany('INSERT OR IGNORE INTO people (id, name) VALUES (?, ?)', to_insert)
            self.person_cache.update(pid for pid, _ in to_insert)
            self.conn.commit()

    def store_movie(self, movie_id, details):
        cursor = self.conn.cursor()
        cursor.execute('''
            INSERT OR REPLACE INTO movies (id, title, year, parental_guide, runtime, genre)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (movie_id, details['title'], details['year'], details['parental_guide'],
              details['runtime'], details['genre']))
        all_relations = []
        all_relations.extend((pid, name, 'movie_directors') for pid, name in details['directors'])
        all_relations.extend((pid, name, 'movie_writers') for pid, name in details['writers'])
        all_relations.extend((pid, name, 'movie_stars') for pid, name in details['stars'][:10])
        if all_relations:
            people = [(pid, name) for pid, name, _ in all_relations]
            self._insert_people(people)
            for pid, _, table in all_relations:
                cursor.execute(f'INSERT OR IGNORE INTO {table} (movie_id, person_id) VALUES (?, ?)', (movie_id, pid))
        if details['gross'] and details['gross'] != 'N/A':
            cursor.execute('INSERT OR REPLACE INTO box_office (movie_id, gross_us_canada) VALUES (?, ?)',
                           (movie_id, details['gross']))
        self.conn.commit()


class IMDBCrawler:
    def __init__(self, excel_path="imdb_dashboard.xlsx", db_name="imdb_movies.db", html_dir="imdb_pages"):
        self.excel_path = excel_path
        self.downloader = IMDBPageDownloader(html_dir)
        self.extractor = IMDBDataExtractor(db_name)

    def run(self):
        try:
            df = pd.read_excel(self.excel_path, sheet_name='Movies')
        except Exception as e:
            print(f"Failed to read Excel: {e}")
            sys.exit(1)
        self.extractor.connect()
        self.downloader.start()
        total = len(df)
        print(f"Processing {total} movies...")
        for idx, row in df.iterrows():
            imdb_id = row.get('imdb_id')
            url = row.get('url')
            if not url or pd.isna(url) or not imdb_id or not str(imdb_id).startswith('tt'):
                print(f"Skipping row {idx+1}: invalid data")
                continue
            movie_id = int(str(imdb_id)[2:])
            print(f"[{idx+1}/{total}] {url}")
            html_path = self.downloader.fetch(url, movie_id)
            if not html_path:
                continue
            details = self.extractor.extract_from_html(html_path)
            self.extractor.store_movie(movie_id, details)
            time.sleep(1)
        self.downloader.stop()
        self.extractor.close()
        print("Crawling and extraction completed.")


class IMDBAnalyzer:
    def __init__(self, excel_path="imdb_dashboard.xlsx"):
        self.df = pd.read_excel(excel_path, sheet_name='Movies')
        self._clean_data()

    def _clean_data(self):
        
        self.df['rating'] = pd.to_numeric(self.df['rating'], errors='coerce')
        self.df['votes'] = pd.to_numeric(self.df['votes'], errors='coerce')
        self.df['duration_min'] = pd.to_numeric(self.df['duration_min'], errors='coerce')
        
        self.df = self.df.dropna(subset=['rating'])

    def summary_stats(self):
        stats = {
            'تعداد فیلم‌ها': len(self.df),
            'میانگین امتیاز': self.df['rating'].mean(),
            'میانه امتیاز': self.df['rating'].median(),
            'انحراف معیار امتیاز': self.df['rating'].std(),
            'حداقل امتیاز': self.df['rating'].min(),
            'حداکثر امتیاز': self.df['rating'].max(),
            'میانگین تعداد رأی': self.df['votes'].mean(),
            'میانگین مدت زمان (دقیقه)': self.df['duration_min'].mean(),
        }
        return stats

    def top_movies_by_rating(self, n=10):
        return self.df.nlargest(n, 'rating')[['title', 'rating', 'votes']]

    def genre_distribution(self):
        genres = self.df['genre'].dropna().str.split(', ')
        all_genres = [g for sublist in genres for g in sublist]
        return pd.Series(all_genres).value_counts().head(10)

    def rating_distribution(self, bins=[0,5,6,7,8,9,10]):
        return pd.cut(self.df['rating'], bins=bins, right=False).value_counts().sort_index()

    def correlation_rating_votes(self):
        return self.df[['rating', 'votes']].corr().loc['rating', 'votes']

    def generate_report(self):
        stats = self.summary_stats()
        print("="*50)
        print("IMDB Top 250 : ")
        print("="*50)
        for key, val in stats.items():
            print(f"{key:.<30} {val:.2f}" if isinstance(val, float) else f"{key:.<30} {val}")
        print("\n Top 10 movies :")
        print(self.top_movies_by_rating(10).to_string(index=False))
        print("\n Top Ganre 10 :")
        print(self.genre_distribution().to_string())
        print("\n Score in range :")
        print(self.rating_distribution().to_string())
        print(f"\n🔗 ضریب همبستگی امتیاز و تعداد رأی: {self.correlation_rating_votes():.3f}")


class ExcelDashboardGenerator:
    def __init__(self, df, output_path="imdb_dashboard_analyzed.xlsx"):
        self.df = df
        self.output_path = output_path
        self.wb = Workbook()
        self.ws_data = self.wb.active
        self.ws_data.title = "Movies"

    def create(self):
        self._write_data()
        self._create_dashboard()
        self.wb.save(self.output_path)
        print(f"Dashboard save in  {self.output_path}  .")

    def _write_data(self):
        for r in dataframe_to_rows(self.df, index=False, header=True):
            self.ws_data.append(r)
        table = Table(displayName="MoviesTable", ref=self.ws_data.dimensions)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        table.tableStyleInfo = style
        self.ws_data.add_table(table)
        for col in self.ws_data.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            self.ws_data.column_dimensions[column].width = adjusted_width

    def _create_dashboard(self):
        ws_dash = self.wb.create_sheet("Dashboard")
        ws_dash.column_dimensions['A'].width = 20
        ws_dash.column_dimensions['B'].width = 20
        ws_dash.column_dimensions['C'].width = 20
        ws_dash.column_dimensions['D'].width = 20
        ws_dash['A1'] = "IMDb Top 250 Dashboard"
        ws_dash['A1'].font = Font(size=16, bold=True)
        ws_dash.merge_cells('A1:D1')

        # 1. Top 10 Movies by Rating
        top10 = self.df.nlargest(10, 'rating')[['title', 'rating']]
        ws_dash['A3'] = "Top 10 Movies by Rating"
        ws_dash['A3'].font = Font(bold=True, size=12)
        for i, (title, rating) in enumerate(top10.values, start=4):
            ws_dash[f'A{i}'] = title
            ws_dash[f'B{i}'] = rating

        chart1 = BarChart()
        chart1.type = "col"
        chart1.title = "Top 10 Movies by Rating"
        chart1.x_axis.title = "Movie"
        chart1.y_axis.title = "Rating"
        chart1.style = 10
        data = Reference(ws_dash, min_col=2, min_row=4, max_row=13)
        cats = Reference(ws_dash, min_col=1, min_row=4, max_row=13)
        chart1.add_data(data, titles_from_data=False)
        chart1.set_categories(cats)
        ws_dash.add_chart(chart1, "D3")

        # 2. Rating distribution
        bins = [0, 5.0, 6.0, 7.0, 8.0, 9.0, 10.0]
        labels = ['<5.0', '5.0-6.0', '6.0-7.0', '7.0-8.0', '8.0-9.0', '9.0+']
        self.df['rating_bin'] = pd.cut(self.df['rating'], bins=bins, labels=labels, right=False)
        hist = self.df['rating_bin'].value_counts().sort_index()
        ws_dash['F3'] = "Rating Distribution"
        ws_dash['F3'].font = Font(bold=True, size=12)
        for i, (bin_label, count) in enumerate(hist.items(), start=4):
            ws_dash[f'F{i}'] = bin_label
            ws_dash[f'G{i}'] = count

        chart2 = BarChart()
        chart2.type = "col"
        chart2.title = "Rating Distribution"
        chart2.x_axis.title = "Rating Range"
        chart2.y_axis.title = "Number of Movies"
        chart2.style = 11
        data = Reference(ws_dash, min_col=7, min_row=4, max_row=4+len(hist)-1)
        cats = Reference(ws_dash, min_col=6, min_row=4, max_row=4+len(hist)-1)
        chart2.add_data(data, titles_from_data=False)
        chart2.set_categories(cats)
        ws_dash.add_chart(chart2, "K3")

        # 3. Top Genres (Pie)
        genre_series = self.df['genre'].dropna().str.split(', ')
        genres = [g for sub in genre_series for g in sub]
        genre_counts = pd.Series(genres).value_counts().head(10)
        ws_dash['F18'] = "Top 10 Genres"
        ws_dash['F18'].font = Font(bold=True, size=12)
        for i, (genre, count) in enumerate(genre_counts.items(), start=19):
            ws_dash[f'F{i}'] = genre
            ws_dash[f'G{i}'] = count

        chart3 = PieChart()
        chart3.title = "Top 10 Genres"
        chart3.style = 2
        data = Reference(ws_dash, min_col=7, min_row=19, max_row=19+len(genre_counts)-1)
        cats = Reference(ws_dash, min_col=6, min_row=19, max_row=19+len(genre_counts)-1)
        chart3.add_data(data, titles_from_data=False)
        chart3.set_categories(cats)
        ws_dash.add_chart(chart3, "K18")

        # 4. Rating vs Votes (Scatter)
        ws_dash['F28'] = "Rating vs Votes"
        ws_dash['F28'].font = Font(bold=True, size=12)
        start_row = 29
        for i, (rating, votes) in enumerate(self.df[['rating', 'votes']].dropna().values, start=start_row):
            ws_dash[f'F{i}'] = rating
            ws_dash[f'G{i}'] = votes

        chart4 = ScatterChart()
        chart4.title = "Rating vs Votes"
        chart4.x_axis.title = "Rating"
        chart4.y_axis.title = "Number of Votes"
        chart4.style = 13
        xvalues = Reference(ws_dash, min_col=6, min_row=start_row, max_row=start_row+len(self.df)-1)
        yvalues = Reference(ws_dash, min_col=7, min_row=start_row, max_row=start_row+len(self.df)-1)
        series = Series(yvalues, xvalues, title_from_data=False)
        chart4.series.append(series)
        ws_dash.add_chart(chart4, "K28")

        # header style
        for row in ws_dash.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
                cell.font = Font(color="FFFFFF", bold=True)
